"""
fix.py  (исправленная версия)
Восстанавливает связь LED <-> пиксели на основе готовых PNG + JSON
"""

from collections import defaultdict
import json
import os
from PIL import Image
import openpyxl

PNG_DIR = "flags_for_letters"
JSON_DIR = "flags_for_letters_json"

colors = defaultdict(lambda: defaultdict(str))      # colors[led_name][png_name] = hex
positions = defaultdict(lambda: defaultdict(set))   # positions[png_name][color] = set of (row, col) 1-based

png_files = [f for f in os.listdir(PNG_DIR) if f.lower().endswith(".png")]
print(f"Найдено PNG-файлов: {len(png_files)}\n")

for png_name in sorted(png_files):
    group = os.path.splitext(png_name)[0]
    json_path = os.path.join(JSON_DIR, f"{group}.json")

    if not os.path.exists(json_path):
        print(f"⚠️  Пропущен {png_name} — нет JSON")
        continue

    # === Загружаем JSON ===
    with open(json_path, encoding="utf-8") as f:
        mapping = json.load(f)

    # === Загружаем PNG и ПРИНУДИТЕЛЬНО переводим в RGB (главное исправление) ===
    img = Image.open(os.path.join(PNG_DIR, png_name)).convert('RGB')
    pixels = img.load()
    width, height = img.size

    print(f"Обрабатываю {png_name} ({width}×{height})")

    # Собираем цвета всех LED для этой картинки
    led_colors = {}
    for letter, data in mapping.items():
        if isinstance(data, dict) and "colors" in data:
            for i, color in enumerate(data["colors"]):
                led_name = f"{letter}_{i+1:03d}"          # FL_001, FR_774, I_1043 ...
                led_colors[led_name] = color.upper()

    # Перебираем все пиксели (точно как в оригинальном main.py)
    for row in range(height):
        for col in range(width):
            r, g, b = pixels[col, row]                    # теперь всегда 3 значения
            hex_color = f"{r:02X}{g:02X}{b:02X}"

            positions[png_name][hex_color].add((row + 1, col + 1))   # 1-based

    # Заполняем colors
    for led_name, hex_color in led_colors.items():
        colors[led_name][png_name] = hex_color

# === Сохраняем результат ===
colors_final = {led: dict(pics) for led, pics in colors.items()}

positions_final = {}
for png_name, color_dict in positions.items():
    positions_final[png_name] = {}
    for color, pos_set in color_dict.items():
        positions_final[png_name][color] = sorted(pos_set)   # list для JSON

with open("colors_map.json", "w", encoding="utf-8") as f:
    json.dump(colors_final, f, indent=2, ensure_ascii=False)

with open("positions_map.json", "w", encoding="utf-8") as f:
    json.dump(positions_final, f, indent=2, ensure_ascii=False)

print("\n✅ ГОТОВО!")
print(f"   colors_map.json    → {len(colors_final)} LED × {len(png_files)} картинок")
print(f"   positions_map.json → {len(positions_final)} картинок")
print("\nТеперь можно анализировать уникальные цвета и восстанавливать таблицу.")



# Пример: найти все LED, у которых цвет уникален хотя бы в одной картинке
unique = {}
for led, pics in colors_final.items():
    for png, color in pics.items():
        if len(positions_final[png].get(color, [])) == 1:
            pos = positions_final[png][color][0]
            unique[led] = (png, color, pos)
            break
print(f"{len(unique)}/{len(colors_final)} LED уже можно точно сопоставить!")



# ====================== СОЗДАНИЕ mapping.json И table.xlsx ======================
print("\n" + "="*70)
print("Восстанавливаем полную карту и создаём table.xlsx...")

# 1. Создаём mapping.json (точно как ты просил)
mapping = {}
for led, pics in colors_final.items():
    for png_name, color in pics.items():
        pos_list = positions_final.get(png_name, {}).get(color, [])
        if len(pos_list) == 1:
            row, col = pos_list[0]
            mapping[led] = [row, col]          # [row, col] 1-based
            break

with open("mapping.json", "w", encoding="utf-8") as f:
    json.dump(mapping, f, indent=2, ensure_ascii=False)

print(f"✓ mapping.json создан — {len(mapping)} LED")

# 2. Создаём table.xlsx (полностью совместимый с оригинальным main.py)
# Берём размер таблицы из первой картинки (как делает main.py)
first_png_name = list(positions_final.keys())[0]
img = Image.open(os.path.join(PNG_DIR, first_png_name))
width, height = img.size

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Mapping"

print(f"Создаём таблицу размером {width} × {height} (как в оригинале)")

# Заполняем таблицу
filled = 0
for led_name, (row, col) in mapping.items():
    ws.cell(row=row, column=col, value=led_name)
    filled += 1

# Немного форматирования (удобно смотреть в Excel)
for col in range(1, width + 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 10

wb.save("table.xlsx")
print(f"✓ table.xlsx создан! Заполнено ячеек: {filled}")

print("\n" + "="*70)
print("ГОТОВО! Теперь у тебя есть:")
print("   • mapping.json          ← простой словарь LED → координаты")
print("   • table.xlsx            ← полноценная таблица для main.py")
print("Можешь запускать main.py — он теперь будет работать с этой таблицей!")